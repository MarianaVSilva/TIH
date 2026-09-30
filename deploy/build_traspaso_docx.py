# -*- coding: utf-8 -*-
"""Genera TRASPASO.docx a partir de TRASPASO.md, editable en Word.

A diferencia del PDF, aqui la prioridad es que se pueda EDITAR comodo:
- Titulos con los estilos integrados de Word (Heading 1/2/3), para que
  funcionen el panel de navegacion y el indice automatico.
- Tablas de Word de verdad, no texto alineado.
- Indice como CAMPO: Word lo recalcula solo al actualizar (F9), asi que
  sigue siendo correcto despues de editar.
"""
import io, os, re
import docx
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ORIGEN = os.path.join(RAIZ, "TRASPASO.md")
DESTINO = os.path.join(RAIZ, "TRASPASO.docx")

VERDE = RGBColor(0x10, 0x31, 0x2B)
VERDE2 = RGBColor(0x1A, 0x4A, 0x3A)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
GRIS = RGBColor(0x47, 0x55, 0x69)
AMBAR = RGBColor(0x92, 0x40, 0x0E)
TINTA = RGBColor(0x1F, 0x24, 0x30)


def sombrear(elem, hexcolor):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hexcolor)
    elem.append(sh)


def sombrear_parrafo(p, hexcolor):
    sombrear(p._p.get_or_add_pPr(), hexcolor)


def borde_izq(p, hexcolor, ancho=18):
    pPr = p._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single"); left.set(qn("w:sz"), str(ancho))
    left.set(qn("w:space"), "6"); left.set(qn("w:color"), hexcolor)
    bd.append(left); pPr.append(bd)


# ── Texto en linea: **negrita**, *cursiva*, `codigo` ──────────────────────────
TOKEN = re.compile(r"(\*\*.+?\*\*(?!\*)|(?<![\w*])\*[^*\n]+\*(?![\w*])|`[^`]+`)")


def _runs(p, texto, negrita=False, cursiva=False):
    """Emite los runs de un fragmento.

    Va por recursion porque el marcado se ANIDA: "**`archivo.md`**" es negrita
    envolviendo codigo. Resolviendo cada token de golpe, el contenido de la
    negrita salia como texto plano y los acentos graves quedaban a la vista.
    """
    for tr in TOKEN.split(texto):
        if not tr:
            continue
        if tr.startswith("**") and tr.endswith("**"):
            _runs(p, tr[2:-2], True, cursiva)
        elif tr.startswith("`") and tr.endswith("`"):
            r = p.add_run(tr[1:-1])
            r.font.name = "Consolas"; r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x0B, 0x3A, 0x30)
            r.bold = negrita; r.italic = cursiva
        elif tr.startswith("*") and tr.endswith("*"):
            _runs(p, tr[1:-1], negrita, True)
        else:
            r = p.add_run(tr); r.bold = negrita; r.italic = cursiva


def escribir(p, texto, base=None):
    texto = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", texto)
    texto = texto.replace("⚠️", "").replace("✅", "").replace("→", "->")
    _runs(p, texto)
    if base:
        for r in p.runs:
            r.font.size = r.font.size or base


doc = Document()

# Margenes y fuente base
for s in doc.sections:
    s.left_margin = s.right_margin = Inches(0.9)
    s.top_margin = s.bottom_margin = Inches(0.9)
n = doc.styles["Normal"]
n.font.name = "Calibri"; n.font.size = Pt(10.5); n.font.color.rgb = TINTA
n.paragraph_format.space_after = Pt(6)

for nom, tam, col in (("Heading 1", 20, VERDE), ("Heading 2", 15, VERDE2), ("Heading 3", 12, TEAL)):
    st = doc.styles[nom]
    st.font.name = "Calibri"; st.font.size = Pt(tam); st.font.bold = True
    st.font.color.rgb = col
    st.paragraph_format.space_before = Pt(14); st.paragraph_format.space_after = Pt(6)

# ── Portada ──────────────────────────────────────────────────────────────────
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("TIH"); r.font.size = Pt(40); r.bold = True; r.font.color.rgb = VERDE
sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Tablero de Información de Hidrocarburos")
r.font.size = Pt(13); r.font.color.rgb = GRIS
doc.add_paragraph()
t2 = doc.add_paragraph(); t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t2.add_run("Documento de traspaso"); r.font.size = Pt(26); r.bold = True; r.font.color.rgb = VERDE
s2 = doc.add_paragraph(); s2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = s2.add_run("Todo lo necesario para recibir y mantener el proyecto")
r.font.size = Pt(12); r.font.color.rgb = GRIS
doc.add_paragraph()

datos = [("Proyecto", "TIH — Tablero de Información de Hidrocarburos"),
         ("En producción", "http://172.16.83.21:8091"),
         ("Repositorio", "172.16.196.12/mvazquez/tih.git"),
         ("Entrega", "Mariana Vázquez (mvazquez)"),
         ("Fecha", "Septiembre de 2026")]
tb = doc.add_table(rows=len(datos), cols=2)
tb.style = "Table Grid"; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, (k, v) in enumerate(datos):
    c0, c1 = tb.rows[i].cells
    c0.width = Inches(1.5); c1.width = Inches(4.2)
    p0 = c0.paragraphs[0]; r0 = p0.add_run(k); r0.bold = True; r0.font.size = Pt(10)
    sombrear(c0._tc.get_or_add_tcPr(), "F1F5F9")
    p1 = c1.paragraphs[0]; escribir(p1, v)
    for rr in p1.runs:
        rr.font.size = Pt(10)

doc.add_paragraph()
av = doc.add_paragraph(); sombrear_parrafo(av, "FEF3C7"); borde_izq(av, "D97706")
r = av.add_run("Este documento no contiene contraseñas ni llaves. ")
r.bold = True; r.font.color.rgb = AMBAR; r.font.size = Pt(10)
r = av.add_run("La sección 3 indica qué credenciales existen, dónde viven y por qué canal solicitarlas.")
r.font.color.rgb = AMBAR; r.font.size = Pt(10)

doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ── Indice como campo (Word lo actualiza con F9) ──────────────────────────────
h = doc.add_paragraph("Contenido", style="Heading 1")
pt = doc.add_paragraph()
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), r'TOC \o "1-3" \h \z \u')
sub_p = OxmlElement("w:p")
sub_r = OxmlElement("w:r")
sub_t = OxmlElement("w:t")
sub_t.text = "Clic derecho > Actualizar campos para generar el índice."
sub_r.append(sub_t); sub_p.append(sub_r); fld.append(sub_p)
pt._p.addnext(fld)

doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ── Cuerpo ───────────────────────────────────────────────────────────────────
md = io.open(ORIGEN, encoding="utf-8").read().splitlines()
i = 0
while i < len(md) and not md[i].startswith("## "):
    i += 1
cuerpo = md[i:]


def es_continuacion(j):
    return (j < len(cuerpo) and cuerpo[j].strip()
            and not re.match(r"^\s*[-*] |^\d+\. |^#{1,6} |^\||^```|^>|^---\s*$", cuerpo[j]))


def tabla(filas):
    ncol = max(len(f) for f in filas)
    filas = [f + [""] * (ncol - len(f)) for f in filas]
    t = doc.add_table(rows=len(filas), cols=ncol)
    t.style = "Table Grid"
    ancho_total = 6.7
    prim = min(len(str(f[0])) for f in filas[1:]) if len(filas) > 1 else 5
    if prim <= 3:
        anchos = [0.5] + [(ancho_total - 0.5) / (ncol - 1)] * (ncol - 1)
    else:
        anchos = [ancho_total / ncol] * ncol
    for ri, fila in enumerate(filas):
        for ci, val in enumerate(fila):
            cel = t.rows[ri].cells[ci]
            cel.width = Inches(anchos[ci])
            p = cel.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
            escribir(p, val)
            for rr in p.runs:
                rr.font.size = Pt(9)
                if ri == 0:
                    rr.bold = True; rr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            sombrear(cel._tc.get_or_add_tcPr(),
                     "10312B" if ri == 0 else ("F8FAFC" if ri % 2 == 0 else "FFFFFF"))
    return t


n_ln = 0
while n_ln < len(cuerpo):
    ln = cuerpo[n_ln]

    if ln.startswith("```"):
        buf = []; n_ln += 1
        while n_ln < len(cuerpo) and not cuerpo[n_ln].startswith("```"):
            buf.append(cuerpo[n_ln]); n_ln += 1
        n_ln += 1
        for k, l in enumerate(buf):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0 if k < len(buf) - 1 else 8)
            p.paragraph_format.space_before = Pt(8 if k == 0 else 0)
            p.paragraph_format.left_indent = Inches(0.12)
            sombrear_parrafo(p, "F6F8F7"); borde_izq(p, "0F766E", 14)
            r = p.add_run(l); r.font.name = "Consolas"; r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x0B, 0x3A, 0x30)
        continue

    if ln.startswith("|"):
        filas = []
        while n_ln < len(cuerpo) and cuerpo[n_ln].startswith("|"):
            f = [c.strip() for c in cuerpo[n_ln].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c or "-") for c in f):
                filas.append(f)
            n_ln += 1
        if filas:
            tabla(filas); doc.add_paragraph().paragraph_format.space_after = Pt(2)
        continue

    if ln.startswith(">"):
        buf = []
        while n_ln < len(cuerpo) and cuerpo[n_ln].startswith(">"):
            buf.append(re.sub(r"^>\s?", "", cuerpo[n_ln])); n_ln += 1
        txt = " ".join(x.strip() for x in buf if x.strip())
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.12)
        es_aviso = "⚠" in " ".join(buf) or txt.lower().startswith(("ojo", "importante"))
        sombrear_parrafo(p, "FEF3C7" if es_aviso else "F8FAFC")
        borde_izq(p, "D97706" if es_aviso else "CBD5E1")
        escribir(p, txt)
        for rr in p.runs:
            rr.font.size = Pt(9.5)
            if es_aviso:
                rr.font.color.rgb = AMBAR
        continue

    if ln.startswith("### "):
        doc.add_paragraph(re.sub(r"[`*]", "", ln[4:]), style="Heading 3"); n_ln += 1; continue
    if ln.startswith("## "):
        doc.add_paragraph(re.sub(r"[`*⚠️]", "", ln[3:]).strip(), style="Heading 2"); n_ln += 1; continue
    if ln.strip() == "---":
        n_ln += 1; continue

    m = re.match(r"^(\s*)[-*] (.+)", ln)
    if m:
        buf = [m.group(2)]; prof = len(m.group(1)) // 2; n_ln += 1
        while es_continuacion(n_ln):
            buf.append(cuerpo[n_ln].strip()); n_ln += 1
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.25 + prof * 0.25)
        escribir(p, " ".join(buf))
        continue

    m = re.match(r"^(\d+)\. (.+)", ln)
    if m:
        buf = [m.group(2)]; n_ln += 1
        while es_continuacion(n_ln):
            buf.append(cuerpo[n_ln].strip()); n_ln += 1
        p = doc.add_paragraph(style="List Number")
        escribir(p, " ".join(buf))
        continue

    if ln.strip():
        buf = [ln]; n_ln += 1
        while es_continuacion(n_ln):
            buf.append(cuerpo[n_ln]); n_ln += 1
        p = doc.add_paragraph()
        escribir(p, " ".join(x.strip() for x in buf))
        continue

    n_ln += 1

doc.save(DESTINO)
print("OK ->", DESTINO)
